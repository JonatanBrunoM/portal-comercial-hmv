"""Pré-validação determinística dos relatórios HMV2670.

Esta etapa não grava dados no Supabase. Ela apenas interpreta o arquivo recebido,
valida estrutura, competência e fechamentos financeiros antes da confirmação.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from typing import Any
import xml.etree.ElementTree as ET

from openpyxl import load_workbook


EXPECTED_HEADERS = {
    "NOME_MEDICO",
    "VALOR_MATERIAL_ESPECIAL",
    "VALOR_TOTAL",
    "SEQ_ORCAMENTO",
    "SOLICITANTE",
    "DATA",
    "VALOR",
    "DESCRICAO_PRO_FAT",
    "SEQ_ORCAMENTO_ITEM",
    "CD_PRO_FAT",
    "QUANTIDADE",
    "VALOR_UNITARIO",
    "VALOR_TOTAL1",
}

MONTHS_PT = {
    1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril",
    5: "Maio", 6: "Junho", 7: "Julho", 8: "Agosto",
    9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro",
}


@dataclass(frozen=True)
class ValidationIssue:
    severity: str
    code: str
    message: str


def _decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return Decimal("0")
    raw = str(value).strip()
    if "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    if not raw:
        return None
    for fmt in ("%d/%m/%y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def _money(value: Decimal) -> str:
    text = f"{value:,.2f}".replace(",", "#").replace(".", ",").replace("#", ".")
    return f"R$ {text}"


def _analyse_records(
    *,
    records: list[dict[str, Any]],
    item_rows: int,
    filename: str,
    digest: str,
    source_format: str,
    source_label: str,
) -> dict[str, Any]:
    issues: list[ValidationIssue] = []
    budget_numbers: set[str] = set()
    repeated_budget_headers: set[str] = set()
    dates: list[date] = []
    invalid_dates = 0
    invalid_financial = 0
    financial_mismatch = 0
    negative_financial = 0
    total_value = Decimal("0")
    procedure_value = Decimal("0")
    material_value = Decimal("0")

    for record in records:
        budget_number = str(record.get("SEQ_ORCAMENTO") or "").strip()
        if not budget_number:
            issues.append(ValidationIssue("CRITICAL", "MISSING_BUDGET_NUMBER", "Há orçamento sem número identificador."))
            continue
        if budget_number in budget_numbers:
            repeated_budget_headers.add(budget_number)
        budget_numbers.add(budget_number)

        parsed_date = _date(record.get("DATA"))
        if parsed_date is None:
            invalid_dates += 1
        else:
            dates.append(parsed_date)

        procedure = _decimal(record.get("VALOR"))
        material = _decimal(record.get("VALOR_MATERIAL_ESPECIAL"))
        total = _decimal(record.get("VALOR_TOTAL"))
        if procedure is None or material is None or total is None:
            invalid_financial += 1
            continue

        procedure_value += procedure
        material_value += material
        total_value += total
        if procedure < 0 or material < 0 or total < 0:
            negative_financial += 1
        if abs((procedure + material) - total) > Decimal("0.01"):
            financial_mismatch += 1

    if not budget_numbers:
        issues.append(ValidationIssue("CRITICAL", "NO_BUDGETS", "Nenhum orçamento foi identificado no arquivo."))
    if invalid_dates:
        issues.append(ValidationIssue("CRITICAL", "INVALID_DATES", f"{invalid_dates} orçamento(s) possuem data ausente ou inválida."))
    if invalid_financial:
        issues.append(ValidationIssue("CRITICAL", "INVALID_FINANCIAL", f"{invalid_financial} orçamento(s) possuem valor financeiro inválido."))
    if financial_mismatch:
        issues.append(ValidationIssue("CRITICAL", "FINANCIAL_MISMATCH", f"{financial_mismatch} orçamento(s) não fecham Procedimento + Material = Total."))
    if repeated_budget_headers:
        issues.append(ValidationIssue("CRITICAL", "DUPLICATE_BUDGET_HEADER", f"{len(repeated_budget_headers)} número(s) de orçamento aparecem mais de uma vez no mesmo arquivo."))
    if negative_financial:
        issues.append(ValidationIssue("WARNING", "NEGATIVE_VALUES", f"{negative_financial} orçamento(s) possuem valor financeiro negativo e exigem conferência."))

    competences = sorted({(d.year, d.month) for d in dates})
    if len(competences) > 1:
        issues.append(ValidationIssue("CRITICAL", "MULTIPLE_COMPETENCES", "O arquivo contém orçamentos de mais de uma competência mensal. A carga deve ser separada por mês."))

    start_date = min(dates) if dates else None
    end_date = max(dates) if dates else None
    span_days = (end_date - start_date).days + 1 if start_date and end_date else 0
    coverage_type = "SEMANAL" if 1 <= span_days <= 8 else "PERÍODO AMPLIADO"
    competence = None
    competence_label = None
    if len(competences) == 1:
        year, month = competences[0]
        competence = f"{year:04d}-{month:02d}-01"
        competence_label = f"{MONTHS_PT[month]}/{year}"

    if not issues:
        issues.append(ValidationIssue("OK", "VALIDATED", "Estrutura, competência e fechamento financeiro passaram pelas validações determinísticas."))

    difference = (procedure_value + material_value) - total_value
    return {
        "filename": filename,
        "sha256": digest,
        "source_format": source_format,
        "source_label": source_label,
        "budget_records": records,
        "valid_for_import": not any(issue.severity == "CRITICAL" for issue in issues),
        "physical_rows": len(records),
        "budgets": len(budget_numbers),
        "item_rows": item_rows,
        "start_date": start_date.isoformat() if start_date else None,
        "end_date": end_date.isoformat() if end_date else None,
        "span_days": span_days,
        "coverage_type": coverage_type,
        "competence": competence,
        "competence_label": competence_label,
        "procedure_value": str(procedure_value),
        "material_value": str(material_value),
        "total_value": str(total_value),
        "procedure_value_label": _money(procedure_value),
        "material_value_label": _money(material_value),
        "total_value_label": _money(total_value),
        "financial_difference": str(difference),
        "financial_difference_label": _money(difference),
        "issues": [asdict(issue) for issue in issues],
    }


def _inspect_xml(content: bytes, filename: str, digest: str) -> dict[str, Any]:
    text = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Não foi possível identificar a codificação do XML.")

    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        raise ValueError(f"XML inválido ou corrompido: {exc}.") from exc

    if root.tag.upper() != "HMV2670":
        raise ValueError(f"Relatório incompatível: raiz XML '{root.tag}' em vez de 'HMV2670'.")

    budget_nodes = root.findall("./LIST_G_SEQ_ORCAMENTO/G_SEQ_ORCAMENTO")
    records: list[dict[str, Any]] = []
    item_rows = 0
    for node in budget_nodes:
        records.append({
            "SEQ_ORCAMENTO": node.findtext("SEQ_ORCAMENTO"),
            "DATA": node.findtext("DATA"),
            "VALOR": node.findtext("VALOR"),
            "VALOR_MATERIAL_ESPECIAL": node.findtext("VALOR_MATERIAL_ESPECIAL"),
            "VALOR_TOTAL": node.findtext("VALOR_TOTAL"),
        })
        item_rows += len(node.findall("./LIST_G_SEQ_ORCAMENTO_ITEM/G_SEQ_ORCAMENTO_ITEM"))

    return _analyse_records(
        records=records,
        item_rows=item_rows,
        filename=filename,
        digest=digest,
        source_format="XML",
        source_label="Oracle Reports · HMV2670",
    )


def _inspect_xlsx(content: bytes, filename: str, digest: str) -> dict[str, Any]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Não foi possível abrir o arquivo como planilha XLSX válida.") from exc

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValueError("A planilha não possui conteúdo.") from exc

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    positions = {name: index for index, name in enumerate(headers) if name}
    missing_headers = sorted(EXPECTED_HEADERS - set(positions))
    if missing_headers:
        return {
            "filename": filename,
            "sha256": digest,
            "source_format": "XLSX",
            "source_label": "Planilha HMV2670",
            "valid_for_import": False,
            "issues": [asdict(ValidationIssue(
                "CRITICAL", "STRUCTURE_MISMATCH",
                "Estrutura incompatível: colunas obrigatórias ausentes: " + ", ".join(missing_headers),
            ))],
        }

    records: list[dict[str, Any]] = []
    item_rows = 0
    for row in rows:
        item_raw = row[positions["SEQ_ORCAMENTO_ITEM"]]
        if item_raw is not None and str(item_raw).strip():
            item_rows += 1
        budget_raw = row[positions["SEQ_ORCAMENTO"]]
        if budget_raw is None or not str(budget_raw).strip():
            continue
        records.append({
            "SEQ_ORCAMENTO": budget_raw,
            "DATA": row[positions["DATA"]],
            "VALOR": row[positions["VALOR"]],
            "VALOR_MATERIAL_ESPECIAL": row[positions["VALOR_MATERIAL_ESPECIAL"]],
            "VALOR_TOTAL": row[positions["VALOR_TOTAL"]],
        })

    return _analyse_records(
        records=records,
        item_rows=item_rows,
        filename=filename,
        digest=digest,
        source_format="XLSX",
        source_label=f"Planilha HMV2670 · {sheet.title}",
    )


def inspect_hmv2670(content: bytes, filename: str) -> dict[str, Any]:
    if not content:
        raise ValueError("O arquivo está vazio.")

    digest = sha256(content).hexdigest()
    lower_name = filename.lower()
    if lower_name.endswith(".xml"):
        return _inspect_xml(content, filename, digest)
    if lower_name.endswith(".xlsx"):
        return _inspect_xlsx(content, filename, digest)

    raise ValueError("Formato não aceito. Envie o relatório HMV2670 original em .XML ou a versão .XLSX.")
