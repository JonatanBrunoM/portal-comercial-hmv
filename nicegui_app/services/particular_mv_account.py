"""Pré-validação do relatório de contas MV (sem acesso ou escrita no banco)."""
from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

HEADERS = (
    "paciente", "atendimento", "tipo_atend", "conta", "convenio_conta",
    "inicio_conta", "final_conta", "conta_fechada", "avisos_concatenados",
    "prestador_atend", "vl_total_conta",
)


def _key(value: Any) -> str:
    value = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    value = "".join(c for c in value if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "_", value).strip("_")


def _text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _integer(value: Any, label: str, row: int) -> int:
    if isinstance(value, bool):
        raise ValueError(f"Linha {row}: {label} inválido.")
    raw = _text(value)
    if not raw:
        raise ValueError(f"Linha {row}: {label} ausente.")
    if isinstance(value, float):
        if not value.is_integer() or abs(value) >= 10**15:
            raise ValueError(f"Linha {row}: {label} perdeu precisão ou não é inteiro.")
        raw = str(int(value))
    elif isinstance(value, Decimal):
        if value != value.to_integral_value():
            raise ValueError(f"Linha {row}: {label} não é inteiro.")
        raw = str(int(value))
    if not re.fullmatch(r"[0-9]+", raw):
        raise ValueError(f"Linha {row}: {label} inválido: {raw!r}.")
    number = int(raw)
    if not 0 < number < 2**63:
        raise ValueError(f"Linha {row}: {label} fora do intervalo.")
    return number


def _date(value: Any, label: str, row: int) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    raw = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Linha {row}: {label} inválida: {raw!r}.")


def _closed(value: Any, row: int) -> bool:
    if isinstance(value, bool):
        return value
    raw = _key(value)
    if raw in {"s", "sim", "1", "true", "fechada", "fechado"}:
        return True
    if raw in {"n", "nao", "0", "false", "aberta", "aberto"}:
        return False
    raise ValueError(f"Linha {row}: situação de conta não reconhecida: {value!r}.")


def _money(value: Any, row: int) -> str:
    if value is None or str(value).strip() == "":
        raise ValueError(f"Linha {row}: valor da conta ausente.")
    raw = str(value).strip().replace("R$", "").replace("\u00a0", "").replace(" ", "")
    if "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    try:
        amount = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"Linha {row}: valor financeiro inválido: {value!r}.") from exc
    if not amount.is_finite() or amount != amount.quantize(Decimal("0.01")):
        raise ValueError(f"Linha {row}: valor financeiro não possui precisão de centavos.")
    if abs(amount) >= Decimal("10000000000000000"):
        raise ValueError(f"Linha {row}: valor financeiro fora do intervalo.")
    return format(amount, ".2f")


def _notices(value: Any, row: int) -> list[str]:
    if value is None or str(value).strip() == "":
        return []
    raw = str(value).strip()
    # Aceita apenas separadores explícitos; não divide sequências numéricas arbitrariamente.
    parts = re.split(r"\s*[,;|\n\r]+\s*", raw)
    notices = []
    for part in parts:
        part = part.strip()
        if not re.fullmatch(r"[0-9]+(?:\.0)?", part):
            raise ValueError(
                f"Linha {row}: formato de avisos não reconhecido ({raw[:100]!r}). "
                "Verifique o separador original antes de importar."
            )
        normalized = part[:-2] if part.endswith(".0") else part
        if normalized not in notices:
            notices.append(normalized)
    return notices


def inspect_mv_accounts(content: bytes, filename: str) -> dict[str, Any]:
    """Analisa XLSX original; devolve payload pronto para revisão, sem gravar."""
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Envie um relatório .xlsx.")
    if not content:
        raise ValueError("O arquivo está vazio.")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = wb.active
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ValueError("A planilha está vazia.")
        normalized = [_key(v) for v in header_row]
        if len(set(v for v in normalized if v)) != len([v for v in normalized if v]):
            raise ValueError("O cabeçalho possui nomes de colunas duplicados.")
        missing = [name for name in HEADERS if name not in normalized]
        if missing:
            raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(missing))
        index = {name: normalized.index(name) for name in HEADERS}
        records: list[dict[str, Any]] = []
        errors: list[str] = []
        account_counts: Counter[int] = Counter()
        attendance_set: set[int] = set()
        closed_count = 0
        notices_count = 0
        total = Decimal("0.00")
        for row_number, cells in enumerate(iterator, start=2):
            if not any(v is not None and str(v).strip() for v in cells):
                continue
            source = {name: cells[i] if i < len(cells) else None for name, i in index.items()}
            try:
                attendance = _integer(source["atendimento"], "atendimento", row_number)
                account = _integer(source["conta"], "conta", row_number)
                closed = _closed(source["conta_fechada"], row_number)
                amount = _money(source["vl_total_conta"], row_number)
                start = _date(source["inicio_conta"], "início da conta", row_number)
                end = _date(source["final_conta"], "final da conta", row_number)
                if start and end and end < start:
                    raise ValueError(f"Linha {row_number}: final da conta anterior ao início.")
                notices = _notices(source["avisos_concatenados"], row_number)
                records.append({
                    "source_row_number": row_number,
                    "attendance_number": attendance,
                    "account_number": account,
                    "patient_name": _text(source["paciente"]),
                    "attendance_type": _text(source["tipo_atend"]),
                    "insurance_name": _text(source["convenio_conta"]),
                    "account_start_date": start,
                    "account_end_date": end,
                    "is_closed": closed,
                    "provider_name": _text(source["prestador_atend"]),
                    "account_total": amount,
                    "notices": notices,
                    "raw_payload": {k: (v.isoformat() if isinstance(v, (date, datetime)) else str(v) if v is not None else None) for k, v in source.items()},
                })
                account_counts[account] += 1
                attendance_set.add(attendance)
                closed_count += int(closed)
                notices_count += len(notices)
                total += Decimal(amount)
            except ValueError as exc:
                errors.append(str(exc))
        duplicates = [number for number, count in account_counts.items() if count > 1]
        if duplicates:
            errors.append(f"Há {len(duplicates)} número(s) de conta repetidos no mesmo arquivo; revisão obrigatória (ex.: {duplicates[:5]}).")
        if not records:
            errors.append("Nenhuma conta válida foi encontrada.")
        return {
            "source_filename": filename,
            "file_sha256": hashlib.sha256(content).hexdigest(),
            "records_total": len(records) + sum(1 for error in errors if error.startswith("Linha ")),
            "records_valid": len(records),
            "distinct_attendances": len(attendance_set),
            "closed_accounts": closed_count,
            "open_accounts": len(records) - closed_count,
            "notices_total": notices_count,
            "account_total_sum": format(total, ".2f"),
            "errors": errors,
            "valid_for_import": not errors,
            "records": records if not errors else [],
        }
    finally:
        wb.close()
